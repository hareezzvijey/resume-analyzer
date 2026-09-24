import io

import fitz
from docx import Document
from fastapi.testclient import TestClient

from backend.main import app


client = TestClient(app)


def create_test_pdf() -> bytes:
    document = fitz.open()

    page = document.new_page()

    page.insert_text(
        (72, 72),
        "John Doe\nPython Developer\nPython SQL Machine Learning"
    )

    pdf_bytes = document.tobytes()

    document.close()

    return pdf_bytes


def create_test_docx() -> bytes:
    document = Document()

    document.add_paragraph("John Doe")
    document.add_paragraph("Python Developer")
    document.add_paragraph(
        "Python SQL Machine Learning"
    )

    buffer = io.BytesIO()

    document.save(buffer)

    return buffer.getvalue()


def test_parse_pdf_resume():
    pdf_content = create_test_pdf()

    response = client.post(
        "/api/resumes/parse",
        files={
            "resume": (
                "resume.pdf",
                pdf_content,
                "application/pdf"
            )
        }
    )

    assert response.status_code == 200

    data = response.json()

    assert data["file_name"] == "resume.pdf"
    assert data["file_type"] == "PDF"
    assert "John Doe" in data["text"]
    assert "Python Developer" in data["text"]


def test_parse_docx_resume():
    docx_content = create_test_docx()

    response = client.post(
        "/api/resumes/parse",
        files={
            "resume": (
                "resume.docx",
                docx_content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        }
    )

    assert response.status_code == 200

    data = response.json()

    assert data["file_name"] == "resume.docx"
    assert data["file_type"] == "DOCX"
    assert "John Doe" in data["text"]


def test_parse_unsupported_file():
    response = client.post(
        "/api/resumes/parse",
        files={
            "resume": (
                "resume.txt",
                b"Some resume content",
                "text/plain"
            )
        }
    )

    assert response.status_code == 400

    data = response.json()

    assert "Unsupported file type" in data["detail"]


def test_parse_empty_file():
    response = client.post(
        "/api/resumes/parse",
        files={
            "resume": (
                "resume.pdf",
                b"",
                "application/pdf"
            )
        }
    )

    assert response.status_code == 400

    data = response.json()

    assert "empty" in data["detail"].lower()


def test_parse_without_file():
    response = client.post(
        "/api/resumes/parse"
    )

    assert response.status_code == 422