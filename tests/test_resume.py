import io

import fitz
import pytest
from docx import Document

from backend.modules.resume.parser import (
    extract_text_from_docx,
    extract_text_from_pdf,
)
from backend.modules.resume.service import process_resume
from backend.modules.resume.validator import (
    validate_file_content,
    validate_file_extension,
)


def create_test_pdf() -> bytes:
    """
    Create a small valid PDF in memory for testing.
    """

    document = fitz.open()

    page = document.new_page()

    page.insert_text(
        (72, 72),
        "John Doe\nPython Developer\nPython SQL Machine Learning"
    )

    pdf_bytes = document.tobytes()

    document.close()

    return pdf_bytes


def create_test_docx() -> bytes:
    """
    Create a small valid DOCX in memory for testing.
    """

    document = Document()

    document.add_paragraph("John Doe")
    document.add_paragraph("Python Developer")
    document.add_paragraph(
        "Python SQL Machine Learning"
    )

    buffer = io.BytesIO()

    document.save(buffer)

    return buffer.getvalue()


def test_validate_pdf_extension():
    file_type = validate_file_extension("resume.pdf")

    assert file_type == "PDF"


def test_validate_docx_extension():
    file_type = validate_file_extension("resume.docx")

    assert file_type == "DOCX"


def test_reject_unsupported_file_type():
    with pytest.raises(ValueError):
        validate_file_extension("resume.txt")


def test_reject_empty_file():
    with pytest.raises(ValueError):
        validate_file_content(b"")


def test_extract_text_from_pdf():
    pdf_content = create_test_pdf()

    text = extract_text_from_pdf(pdf_content)

    assert "John Doe" in text
    assert "Python Developer" in text
    assert "Machine Learning" in text


def test_extract_text_from_docx():
    docx_content = create_test_docx()

    text = extract_text_from_docx(docx_content)

    assert "John Doe" in text
    assert "Python Developer" in text
    assert "Machine Learning" in text


def test_process_pdf_resume():
    pdf_content = create_test_pdf()

    result = process_resume(
        "resume.pdf",
        pdf_content
    )

    assert result["file_name"] == "resume.pdf"
    assert result["file_type"] == "PDF"
    assert "John Doe" in result["text"]


def test_process_docx_resume():
    docx_content = create_test_docx()

    result = process_resume(
        "resume.docx",
        docx_content
    )

    assert result["file_name"] == "resume.docx"
    assert result["file_type"] == "DOCX"
    assert "John Doe" in result["text"]


def test_reject_empty_resume():
    with pytest.raises(ValueError):
        process_resume(
            "resume.pdf",
            b""
        )


def test_reject_unsupported_resume():
    with pytest.raises(ValueError):
        process_resume(
            "resume.txt",
            b"some content"
        )